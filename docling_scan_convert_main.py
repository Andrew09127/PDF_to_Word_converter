"""
docling_scan_convert.py
───────────────────────
Универсальный конвертер отсканированных PDF → редактируемый DOCX.

Стек:
  • Docling (DocLayNet layout + TableFormer tables + EasyOCR) — понимание макета
  • python-docx — сборка DOCX с оформлением

Установка:
  pip install docling python-docx
"""

from __future__ import annotations

import gc
import logging
import shutil
import time
from io import BytesIO
from pathlib import Path

from docling.backend.pypdfium2_backend import PyPdfiumDocumentBackend
from docling.datamodel.base_models import InputFormat
from docling.datamodel.pipeline_options import (
    EasyOcrOptions,
    PdfPipelineOptions,
    TableFormerMode,
)
from docling.document_converter import DocumentConverter, PdfFormatOption
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

# ─── Logging ──────────────────────────────────────────────────────────────────
logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s - %(levelname)s - %(message)s",
    handlers=[
        logging.FileHandler("docling_conversion.log", encoding="utf-8"),
        logging.StreamHandler(),
    ],
)
log = logging.getLogger(__name__)

# ─── Page / font constants ─────────────────────────────────────────────────────
PAGE_W_INCH  = 8.27
PAGE_H_INCH  = 11.69
MARGIN_INCH  = 1.0
EMU_PER_INCH = 914400
FONT_NAME    = "Times New Roman"
BODY_PT      = 11.0

# DocItemLabel → font size (pt)
_LABEL_PT: dict[str, float] = {
    "title":          18.0,
    "section_header": 13.0,
    "paragraph":      11.0,
    "text":           11.0,
    "list_item":      11.0,
    "caption":         9.0,
    "footnote":        8.5,
    "page_header":     8.0,
    "page_footer":     8.0,
    "formula":        10.0,
    "code":           10.0,
}

# DocItemLabel → Word heading level
_LABEL_HEADING: dict[str, int] = {
    "title":          1,
    "section_header": 2,
}

_SKIP_LABELS: frozenset[str] = frozenset({"page_header", "page_footer"})


# ══════════════════════════════════════════════════════════════════════════════
#  Docling pipeline
# ══════════════════════════════════════════════════════════════════════════════

def _build_converter() -> DocumentConverter:
    """Configure Docling for scanned (image-only) PDFs."""
    ocr_opts = EasyOcrOptions(lang=["ru", "en"], force_full_page_ocr=True)

    pipeline_opts = PdfPipelineOptions()
    pipeline_opts.do_ocr                                   = True
    pipeline_opts.ocr_options                              = ocr_opts
    pipeline_opts.do_table_structure                       = True
    pipeline_opts.table_structure_options.do_cell_matching = True
    pipeline_opts.table_structure_options.mode             = TableFormerMode.ACCURATE
    pipeline_opts.generate_page_images                     = True

    return DocumentConverter(
        format_options={
            InputFormat.PDF: PdfFormatOption(
                pipeline_options=pipeline_opts,
                backend=PyPdfiumDocumentBackend,
            )
        }
    )


# ══════════════════════════════════════════════════════════════════════════════
#  DOCX helpers
# ══════════════════════════════════════════════════════════════════════════════

def _init_document() -> Document:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width  = int(PAGE_W_INCH * EMU_PER_INCH)
    sec.page_height = int(PAGE_H_INCH * EMU_PER_INCH)
    for attr in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, attr, int(MARGIN_INCH * EMU_PER_INCH))

    normal           = doc.styles["Normal"]
    normal.font.name = FONT_NAME
    normal.font.size = Pt(BODY_PT)

    for level in (1, 2, 3):
        h                = doc.styles[f"Heading {level}"]
        h.font.name      = FONT_NAME
        h.font.color.rgb = RGBColor(0, 0, 0)
        h.font.bold      = True

    return doc


def _set_cell_borders(cell) -> None:
    tc      = cell._tc
    tcPr    = tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "000000")
        borders.append(el)
    tcPr.append(borders)


def _detect_alignment(bbox, page_width: float) -> WD_ALIGN_PARAGRAPH:
    """Map bbox position relative to page width to a Word alignment."""
    if page_width <= 0:
        return WD_ALIGN_PARAGRAPH.LEFT

    x0      = float(getattr(bbox, "l", getattr(bbox, "x0", 0)))
    x1      = float(getattr(bbox, "r", getattr(bbox, "x1", page_width)))
    block_w = x1 - x0
    cx      = (x0 + x1) / 2
    ratio   = block_w / page_width

    if ratio > 0.75:
        return WD_ALIGN_PARAGRAPH.JUSTIFY
    if abs(cx - page_width / 2) < page_width * 0.08 and ratio < 0.68:
        return WD_ALIGN_PARAGRAPH.CENTER
    if x0 > page_width * 0.12 and x1 > page_width * 0.72 and ratio < 0.82:
        return WD_ALIGN_PARAGRAPH.RIGHT
    return WD_ALIGN_PARAGRAPH.LEFT


def _add_heading(doc: Document, text: str, level: int, font_pt: float) -> None:
    para           = doc.add_heading("", level=min(level, 9))
    para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run                = para.add_run(text)
    run.font.name      = FONT_NAME
    run.font.size      = Pt(font_pt)
    run.font.bold      = True
    run.font.color.rgb = RGBColor(0, 0, 0)


def _add_table_from_grid(doc: Document, grid: list) -> None:
    """Build a Word table from Docling's TableData.grid (List[List[TableCell|None]])."""
    if not grid:
        return
    n_rows = len(grid)
    n_cols = max((len(row) for row in grid), default=0)
    if n_cols == 0:
        return

    tbl       = doc.add_table(rows=n_rows, cols=n_cols)
    tbl.style = "Table Grid"

    for r_idx, row in enumerate(grid):
        for c_idx in range(n_cols):
            cell      = tbl.rows[r_idx].cells[c_idx]
            cell_data = row[c_idx] if c_idx < len(row) else None
            text      = (getattr(cell_data, "text", "") or "") if cell_data else ""
            _set_cell_borders(cell)
            for p in cell.paragraphs:
                p.clear()
            para = cell.paragraphs[0]
            para.paragraph_format.space_before = Pt(0)
            para.paragraph_format.space_after  = Pt(0)
            run           = para.add_run(text)
            run.font.name = FONT_NAME
            run.font.size = Pt(9.0)
            run.bold      = (r_idx == 0)

    doc.add_paragraph()


def _add_table_from_cells(doc: Document, table_item) -> None:
    """Fallback: reconstruct grid from table_item.data.table_cells."""
    data   = table_item.data
    n_rows = getattr(data, "num_rows", 0)
    n_cols = getattr(data, "num_cols", 0)
    if n_rows == 0 or n_cols == 0:
        return

    grid: list[list[str]] = [[""] * n_cols for _ in range(n_rows)]
    for cell in getattr(data, "table_cells", []):
        r = getattr(cell, "start_row_offset_idx", 0)
        c = getattr(cell, "start_col_offset_idx", 0)
        if 0 <= r < n_rows and 0 <= c < n_cols:
            grid[r][c] = getattr(cell, "text", "") or ""

    tbl       = doc.add_table(rows=n_rows, cols=n_cols)
    tbl.style = "Table Grid"
    for r_idx, row in enumerate(grid):
        for c_idx, text in enumerate(row):
            cell = tbl.rows[r_idx].cells[c_idx]
            _set_cell_borders(cell)
            for p in cell.paragraphs:
                p.clear()
            para = cell.paragraphs[0]
            para.paragraph_format.space_before = Pt(0)
            para.paragraph_format.space_after  = Pt(0)
            run           = para.add_run(text)
            run.font.name = FONT_NAME
            run.font.size = Pt(9.0)
            run.bold      = (r_idx == 0)
    doc.add_paragraph()


# ══════════════════════════════════════════════════════════════════════════════
#  DoclingDocument → python-docx
# ══════════════════════════════════════════════════════════════════════════════

def _build_docx(dl_doc, page_sizes: dict[int, tuple[float, float]]) -> Document:
    """
    Iterate DoclingDocument in reading order and populate a Word document.

    Inserts a page break whenever item.prov[0].page_no increments, so each
    source PDF page maps to a separate Word page.

    page_sizes: {page_no (1-indexed): (width_pt, height_pt)}
    """
    doc          = _init_document()
    current_page = -1          # tracks last seen page number
    has_content  = False       # guards against a leading page break

    for item, _level in dl_doc.iterate_items():
        raw_label = getattr(item, "label", None)
        if raw_label is None:
            continue
        label_str: str = (
            raw_label.value if hasattr(raw_label, "value") else str(raw_label)
        ).lower()

        if label_str in _SKIP_LABELS:
            continue

        # ── Page-break on page transition ────────────────────────────────────
        prov_list = getattr(item, "prov", None) or []
        if prov_list:
            item_page = int(getattr(prov_list[0], "page_no", current_page))
            if has_content and item_page > current_page:
                doc.add_page_break()
            current_page = item_page

        # ── Tables ──────────────────────────────────────────────────────────
        if label_str == "table":
            try:
                data = getattr(item, "data", None)
                if data is not None and hasattr(data, "grid"):
                    _add_table_from_grid(doc, data.grid)
                else:
                    _add_table_from_cells(doc, item)
                has_content = True
            except Exception as exc:
                log.warning("Table render skipped: %s", exc)
            continue

        # ── Pictures / Figures ──────────────────────────────────────────────
        if label_str in ("picture", "figure", "image"):
            try:
                pil_img = item.get_image(dl_doc)
                if pil_img is not None:
                    buf = BytesIO()
                    pil_img.save(buf, format="PNG")
                    buf.seek(0)
                    para           = doc.add_paragraph()
                    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    para.add_run().add_picture(buf, width=Inches(5.0))
                    has_content = True
            except Exception as exc:
                log.debug("Picture skipped: %s", exc)
            continue

        # ── Text items ──────────────────────────────────────────────────────
        text = (getattr(item, "text", None) or "").strip()
        if not text:
            continue

        font_pt = _LABEL_PT.get(label_str, BODY_PT)

        if label_str in _LABEL_HEADING:
            _add_heading(doc, text, _LABEL_HEADING[label_str], font_pt)
            has_content = True
            continue

        if label_str == "list_item":
            try:
                para = doc.add_paragraph(style="List Bullet")
            except KeyError:
                para = doc.add_paragraph()
            para.paragraph_format.space_before = Pt(0)
            para.paragraph_format.space_after  = Pt(1)
            run           = para.add_run(text)
            run.font.name = FONT_NAME
            run.font.size = Pt(font_pt)
            has_content = True
            continue

        # Alignment from bbox provenance
        alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if prov_list:
            prov    = prov_list[0]
            page_no = getattr(prov, "page_no", 1)
            pw, _   = page_sizes.get(int(page_no), (0.0, 0.0))
            bbox    = getattr(prov, "bbox", None)
            if pw > 0 and bbox is not None:
                alignment = _detect_alignment(bbox, pw)

        bold   = label_str in ("title", "section_header")
        italic = label_str in ("caption", "footnote")

        para = doc.add_paragraph()
        para.alignment                     = alignment
        para.paragraph_format.space_before = Pt(0)
        para.paragraph_format.space_after  = Pt(2.0)
        run           = para.add_run(text)
        run.font.name = FONT_NAME
        run.font.size = Pt(font_pt)
        run.bold      = bold
        run.italic    = italic
        has_content   = True

    return doc


# ══════════════════════════════════════════════════════════════════════════════
#  Single-file entry point
# ══════════════════════════════════════════════════════════════════════════════

def convert_pdf(
    pdf_path: Path,
    docx_path: Path,
    converter: DocumentConverter,
) -> bool:
    log.info("  Конвертация: %s", pdf_path.name)
    try:
        result = converter.convert(str(pdf_path))
        dl_doc = result.document

        page_sizes: dict[int, tuple[float, float]] = {}
        for page_no, page in dl_doc.pages.items():
            size = getattr(page, "size", None)
            if size:
                page_sizes[int(page_no)] = (
                    float(getattr(size, "width",  0.0)),
                    float(getattr(size, "height", 0.0)),
                )

        doc = _build_docx(dl_doc, page_sizes)
        doc.save(str(docx_path))
        log.info("  ✓ %s", docx_path.name)
        return True
    except Exception as exc:
        log.error("  ✗ %s: %s", pdf_path.name, exc, exc_info=True)
        return False


# ══════════════════════════════════════════════════════════════════════════════
#  Batch processor
# ══════════════════════════════════════════════════════════════════════════════

class DoclingBatchConverter:
    def __init__(
        self,
        input_folder: str,
        output_folder: str,
        backup_folder: str | None = None,
    ) -> None:
        self.input_folder  = Path(input_folder)
        self.output_folder = Path(output_folder)
        self.backup_folder = Path(backup_folder) if backup_folder else None
        self.output_folder.mkdir(parents=True, exist_ok=True)
        if self.backup_folder:
            self.backup_folder.mkdir(parents=True, exist_ok=True)
        self._converter: DocumentConverter | None = None

    @property
    def converter(self) -> DocumentConverter:
        if self._converter is None:
            log.info(
                "Инициализация Docling pipeline "
                "(DocLayNet + TableFormer + EasyOCR)..."
            )
            self._converter = _build_converter()
            log.info("Pipeline готов.")
        return self._converter

    def process(self, move_to_backup: bool = True) -> None:
        pdf_files = sorted(self.input_folder.glob("*.pdf"))
        total     = len(pdf_files)
        log.info("Найдено %d PDF-файлов.", total)

        stats = {"ok": 0, "fail": 0, "skip": 0}
        t0    = time.monotonic()

        for idx, pdf_path in enumerate(pdf_files, 1):
            docx_path = self.output_folder / f"{pdf_path.stem}.docx"
            if docx_path.exists():
                log.info("[%d/%d] Пропуск: %s", idx, total, pdf_path.name)
                stats["skip"] += 1
                continue

            log.info(
                "[%d/%d] %s (%.2f MB)",
                idx, total, pdf_path.name,
                pdf_path.stat().st_size / 1_048_576,
            )
            ok = convert_pdf(pdf_path, docx_path, self.converter)

            if ok:
                stats["ok"] += 1
                if self.backup_folder and move_to_backup:
                    shutil.move(
                        str(pdf_path),
                        str(self.backup_folder / pdf_path.name),
                    )
            else:
                stats["fail"] += 1

            if idx % 10 == 0:
                gc.collect()

        elapsed = time.monotonic() - t0
        sep = "=" * 55
        log.info("\n%s", sep)
        log.info(
            "Готово за %.1f мин | OK: %d | Ошибок: %d | Пропущено: %d",
            elapsed / 60, stats["ok"], stats["fail"], stats["skip"],
        )
        log.info(sep)


# ══════════════════════════════════════════════════════════════════════════════
#  Entry point
# ══════════════════════════════════════════════════════════════════════════════

if __name__ == "__main__":
    DoclingBatchConverter(
        input_folder  = "./pdf_to_convert",
        output_folder = "./converted_docling",
        backup_folder = "./pdf_backup",
    ).process(move_to_backup=True)
